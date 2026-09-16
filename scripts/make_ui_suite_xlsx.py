"""Build docs/evaluation/ui-suite.xlsx — an uploadable evaluation suite.

The desktop Evaluation view imports the FIRST sheet of an XLSX (or CSV/TSV/JSON)
with columns: id, goal, working_directory, expected_output_contains, checks,
metadata. The `checks` column accepts a preset name ("ui_bundle") or a
"+"-joined subset ("code+info+pacing+script"); the desktop expands presets into
the same named checks the CLI harness uses (ui_artifact_checks).

Run with: uv run --with openpyxl python scripts/make_ui_suite_xlsx.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parents[1] / "docs" / "evaluation" / "ui-suite.xlsx"

BUNDLE_GOAL = (
    "Produce a complete {name} UI bundle. Structure the answer with these "
    "labeled sections, each introduced by its exact name: \"Component Code\" - "
    "the implementation in a fenced code block, in any framework or plain "
    "HTML; \"Component Info\" - document the inputs, the states, and the "
    "accessibility behaviour; \"Pacing\" - a numbered reveal timeline with "
    "every duration written in ms; \"Script\" - the narration script inside a "
    "fenced code block."
)

ROWS = [
    (
        "card_bundle",
        BUNDLE_GOAL.format(name="interactive-card"),
        "",
        "",
        "ui_bundle",
        "",
    ),
    (
        "navbar_bundle",
        BUNDLE_GOAL.format(name="top-navigation-bar with dropdown menus"),
        "",
        "",
        "ui_bundle",
        "",
    ),
    (
        "pricing_card_bundle",
        BUNDLE_GOAL.format(name="pricing-card with a monthly/yearly toggle"),
        "",
        "",
        "ui_bundle",
        "",
    ),
    (
        "data_table_framework_free",
        (
            "Generate a data-table UI component in any framework of your "
            "choice. Include a \"Component Code\" section with the "
            "implementation in a fenced code block, and a \"Component Info\" "
            "section that documents the inputs and the states."
        ),
        "",
        "",
        "code+info",
        "",
    ),
    (
        "settings_panel_documented",
        (
            "Build a settings-panel UI component and document it fully. "
            "Include a \"Component Code\" section with the implementation in a "
            "fenced code block (any framework), and a \"Component Info\" "
            "section covering inputs, states, and accessibility."
        ),
        "",
        "",
        "code+info",
        "",
    ),
    (
        "hero_section_code_script",
        (
            "Create a landing-page hero section in any framework. Include a "
            "\"Component Code\" section with the implementation in a fenced "
            "code block, and a \"Script\" section with the narration text "
            "inside a fenced code block."
        ),
        "",
        "",
        "code+script",
        "",
    ),
    (
        "onboarding_pacing_script",
        (
            "Plan the presentation layer for a three-step onboarding reveal. "
            "Provide a \"Pacing\" section with a numbered timeline where every "
            "duration is written in ms, and a \"Script\" section with the "
            "narration text inside a fenced code block."
        ),
        "",
        "",
        "pacing+script",
        "",
    ),
    (
        "legacy_readme_summary",
        (
            "Read README.md in the workspace and reply with the project's "
            "name."
        ),
        "",
        "OperatingAgent",
        "",
        "",
    ),
]

README_LINES = [
    ("UI evaluation suite — how to use this file", ""),
    ("", ""),
    ("Upload", "Evaluation view -> 'Import JSON/CSV/TSV/XLSX' -> pick this file -> 'Compare agents'."),
    ("", "Only the first sheet ('cases') is imported; this sheet is ignored."),
    ("", ""),
    ("Columns", ""),
    ("id", "Unique case id shown in the report."),
    ("goal", "The prompt both agents receive. Bundle goals ask for labeled sections."),
    ("working_directory", "Optional. Blank = the folder chosen in the Evaluation view."),
    ("expected_output_contains", "Optional legacy check: output must contain this text."),
    ("checks", "Optional named checks: 'ui_bundle', or '+'-joined: code+info+pacing+script."),
    ("metadata", "Optional JSON object for extra run metadata."),
    ("", ""),
    ("Check presets", ""),
    ("code", "A fenced code block after a 'Component Code' label (any language/framework)."),
    ("info", "After a 'Component Info' label: inputs, states, accessibility documented."),
    ("pacing", "After a 'Pacing' label: at least one explicit duration like '250 ms'."),
    ("script", "A fenced block after a 'Script' label (the narration)."),
    ("ui_bundle", "All four checks — the complete artifact bundle."),
    ("", ""),
    ("Notes", "A case passes only when the run completed and every check passed."),
    ("", "Add rows freely; each row with an id and a goal becomes a case."),
]

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build() -> None:
    workbook = Workbook()
    cases = workbook.active
    cases.title = "cases"
    headers = ["id", "goal", "working_directory", "expected_output_contains", "checks", "metadata"]
    cases.append(headers)
    for cell in cases[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in ROWS:
        cases.append(list(row))
    widths = {"A": 30, "B": 95, "C": 20, "D": 26, "E": 22, "F": 20}
    for column, width in widths.items():
        cases.column_dimensions[column].width = width
    for row in cases.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    cases.freeze_panes = "A2"

    guide = workbook.create_sheet("read me")
    for label, text in README_LINES:
        guide.append([label, text])
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 100
    for row in guide.iter_rows(min_row=1):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    # Header line spans both columns.
    guide.merge_cells("A1:B1")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)
    print(f"wrote {OUT} ({len(ROWS)} cases)")


if __name__ == "__main__":
    build()
